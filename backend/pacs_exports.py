import math
import os
import tempfile
import textwrap
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


def _now_stamp():
    return datetime.utcnow().strftime("%Y%m%d_%H%M%S")


def _bool_from_value(value, default=False):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def _safe_filename(value):
    return "".join(character if character.isalnum() or character in {"-", "_"} else "_" for character in value).strip("_") or "pacs_export"


def _wrap_lines(text, width):
    return textwrap.wrap(text, width=width, break_long_words=False, replace_whitespace=False) or [text]


def _load_font(size=22):
    try:
        from PIL import ImageFont

        return ImageFont.truetype("arial.ttf", size=size)
    except Exception:
        return ImageFont.load_default()


def _text_size(draw, text, font):
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def _draw_box(draw, box, title, lines, fill, outline, font, title_font=None):
    left, top, right, bottom = box
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=4)
    title_font = title_font or font

    title_lines = _wrap_lines(title, 20)
    title_height = sum(_text_size(draw, line, title_font)[1] for line in title_lines) + (len(title_lines) - 1) * 2
    body_lines = []
    for line in lines:
        body_lines.extend(_wrap_lines(line, 28))
    body_height = sum(_text_size(draw, line, font)[1] for line in body_lines) + max(0, len(body_lines) - 1) * 3

    current_y = top + max(16, (bottom - top - title_height - body_height) // 2)
    for line in title_lines:
        width, height = _text_size(draw, line, title_font)
        draw.text(((left + right - width) / 2, current_y), line, font=title_font, fill="#0b1f33")
        current_y += height + 2

    current_y += 10
    for line in body_lines:
        width, height = _text_size(draw, line, font)
        draw.text(((left + right - width) / 2, current_y), line, font=font, fill="#1f2937")
        current_y += height + 3


def _draw_arrow(draw, start, end, color="#334155", width=5):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    arrow_length = 16
    arrow_angle = math.pi / 7
    left = (
        end[0] - arrow_length * math.cos(angle - arrow_angle),
        end[1] - arrow_length * math.sin(angle - arrow_angle),
    )
    right = (
        end[0] - arrow_length * math.cos(angle + arrow_angle),
        end[1] - arrow_length * math.sin(angle + arrow_angle),
    )
    draw.polygon([end, left, right], fill=color)


def _make_base_canvas(title, subtitle):
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    title_font = _load_font(28)
    body_font = _load_font(18)
    title_width, title_height = _text_size(draw, title, title_font)
    draw.text(((1600 - title_width) / 2, 28), title, font=title_font, fill="#0f172a")
    subtitle_width, subtitle_height = _text_size(draw, subtitle, body_font)
    draw.text(((1600 - subtitle_width) / 2, 72), subtitle, font=body_font, fill="#475569")
    return image, draw, title_font, body_font


def _workflow_diagram(output_path):
    image, draw, title_font, body_font = _make_base_canvas(
        "PACS Continuity Workflow",
        "Read-only discovery, metadata preview, local indexing, and export generation",
    )

    source_box = (70, 220, 360, 420)
    index_box = (430, 220, 720, 420)
    registry_box = (790, 220, 1080, 420)
    export_box = (1150, 220, 1530, 420)

    _draw_box(draw, source_box, "Sources", ["FHIR", "DICOMweb", "NAS", "Firebird / RIS shares"], "#dbeafe", "#2563eb", body_font)
    _draw_box(draw, index_box, "Read-only indexer", ["pydicom preview", "identity matching", "source classification"], "#dcfce7", "#16a34a", body_font)
    _draw_box(draw, registry_box, "Local registry", ["SQLite continuity store", "search", "timeline", "audit trail"], "#fef3c7", "#d97706", body_font)
    _draw_box(draw, export_box, "Exports", ["DOCX", "PDF", "redacted report"], "#fce7f3", "#db2777", body_font)

    _draw_arrow(draw, (360, 320), (430, 320))
    _draw_arrow(draw, (720, 320), (790, 320))
    _draw_arrow(draw, (1080, 320), (1150, 320))

    draw.text((92, 470), "Nothing in this flow writes back to source PACS, NAS, or study databases.", font=body_font, fill="#0f172a")
    image.save(output_path)


def _database_diagram(output_path):
    image, draw, title_font, body_font = _make_base_canvas(
        "Local Registry Database Diagram",
        "Simplified schema for search, linkage, and audit reporting",
    )

    registry_sources = (580, 160, 1020, 290)
    empi_patients = (120, 340, 460, 470)
    source_identity_links = (560, 340, 1040, 470)
    imaging_studies = (1120, 340, 1480, 470)
    storage_assets = (260, 580, 620, 710)
    pacs_audit_logs = (980, 580, 1340, 710)

    _draw_box(draw, registry_sources, "registry_sources", ["source_type", "source_label", "source_location", "last_status"], "#dbeafe", "#2563eb", body_font)
    _draw_box(draw, empi_patients, "empi_patients", ["linked patient identity", "EMPI, name, birth date"], "#dcfce7", "#16a34a", body_font)
    _draw_box(draw, source_identity_links, "source_identity_links", ["source record to EMPI map", "source_patient_key", "match_rule"], "#fef3c7", "#d97706", body_font)
    _draw_box(draw, imaging_studies, "imaging_studies", ["study UID", "modality", "date", "description"], "#ede9fe", "#7c3aed", body_font)
    _draw_box(draw, storage_assets, "storage_assets", ["file path", "asset family", "database or DICOM archive"], "#fee2e2", "#ef4444", body_font)
    _draw_box(draw, pacs_audit_logs, "pacs_audit_logs", ["append-only route log", "user, client, model provenance"], "#e2e8f0", "#475569", body_font)

    _draw_arrow(draw, (800, 290), (800, 340))
    _draw_arrow(draw, (460, 405), (560, 405))
    _draw_arrow(draw, (1040, 405), (1120, 405))
    _draw_arrow(draw, (800, 470), (430, 580))
    _draw_arrow(draw, (930, 470), (1120, 580))
    _draw_arrow(draw, (800, 290), (1260, 580), color="#64748b")

    draw.text((92, 760), "The source systems remain untouched; the registry stores only local linkage and metadata.", font=body_font, fill="#0f172a")
    image.save(output_path)


def _network_diagram(output_path):
    image, draw, title_font, body_font = _make_base_canvas(
        "PACS Mentor Network View",
        "How the mentor talks to local and remote systems without editing their source data",
    )

    mentor_ui = (90, 180, 390, 310)
    mentor_api = (490, 170, 900, 330)
    local_registry = (520, 400, 880, 540)
    fhir = (1110, 140, 1510, 250)
    dicomweb = (1110, 290, 1510, 400)
    nas = (1110, 440, 1510, 550)
    audit = (1110, 590, 1510, 700)

    _draw_box(draw, mentor_ui, "PACS Mentor UI", ["chat", "report download", "progress bar"], "#dbeafe", "#2563eb", body_font)
    _draw_box(draw, mentor_api, "Flask PACS endpoint", ["read-only routing", "context-aware guidance", "export generation"], "#dcfce7", "#16a34a", body_font)
    _draw_box(draw, local_registry, "Local SQLite registry", ["search", "timeline", "diagram source data"], "#fef3c7", "#d97706", body_font)
    _draw_box(draw, fhir, "FHIR server", ["patient / worklist metadata"], "#ede9fe", "#7c3aed", body_font)
    _draw_box(draw, dicomweb, "DICOMweb server", ["study metadata", "QIDO-RS"], "#ede9fe", "#7c3aed", body_font)
    _draw_box(draw, nas, "Mounted NAS share", ["DICOM files", "Firebird / RIS folders"], "#fee2e2", "#ef4444", body_font)
    _draw_box(draw, audit, "Audit log", ["append-only", "separate from chat deletion"], "#e2e8f0", "#475569", body_font)

    _draw_arrow(draw, (390, 245), (490, 245))
    _draw_arrow(draw, (695, 330), (695, 400))
    _draw_arrow(draw, (900, 230), (1110, 195), color="#7c3aed")
    _draw_arrow(draw, (900, 240), (1110, 340), color="#7c3aed")
    _draw_arrow(draw, (900, 245), (1110, 495), color="#ef4444")
    _draw_arrow(draw, (900, 285), (1110, 640), color="#475569")

    draw.text((94, 750), "All paths are advisory and read-only: the mentor indexes or previews, but does not rewrite external systems.", font=body_font, fill="#0f172a")
    image.save(output_path)


def _create_diagram_set(temp_dir):
    diagrams = []
    for name, builder in [
        ("workflow", _workflow_diagram),
        ("database", _database_diagram),
        ("network", _network_diagram),
    ]:
        path = os.path.join(temp_dir, f"{name}.png")
        builder(path)
        diagrams.append({"name": name, "path": path})
    return diagrams


def _build_export_context(registry, mentor=None, query=None, redact=True):
    snapshot = registry.registry_snapshot()
    overview = registry.describe_registry()
    search_result = None
    if query:
        search_result = registry.search_registry(query, limit=10, redact=redact)

    summary_lines = [
        f"Generated: {datetime.utcnow().isoformat(timespec='seconds')}Z",
        f"Sources discovered: {snapshot.get('source_count', 0)}",
        f"Linked patients: {snapshot.get('patient_count', 0)}",
        f"Indexed studies: {snapshot.get('study_count', 0)}",
        f"Stored assets: {snapshot.get('asset_count', 0)}",
        "Source systems are not modified by this export.",
    ]

    if mentor and getattr(mentor, 'last_index_progress', None) is not None:
        summary_lines.append(f"Last sync progress: {mentor.last_index_progress}% at {mentor.last_index_step or 'complete'}")

    return {
        "snapshot": snapshot,
        "overview": overview,
        "search_result": search_result,
        "summary_lines": summary_lines,
    }


def _add_docx_paragraphs(document, lines, style_name=None):
    for line in lines:
        paragraph = document.add_paragraph(style=style_name) if style_name else document.add_paragraph()
        paragraph.add_run(line)


def _render_docx(export_context, diagrams, output_path, include_diagrams=True):
    document = Document()
    title = document.add_heading("PACS Mentor Export", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph("Read-only registry summary, workflow diagrams, and export-ready PACS metadata")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle.runs:
        subtitle.runs[0].italic = True

    for line in export_context["summary_lines"]:
        document.add_paragraph(line)

    document.add_heading("Registry Overview", level=1)
    overview = export_context["overview"].split(". ")
    _add_docx_paragraphs(document, [sentence.strip() for sentence in overview if sentence.strip()])

    snapshot = export_context["snapshot"]
    document.add_heading("Source Inventory", level=1)
    for source in snapshot.get("sources", []):
        document.add_paragraph(
            f"{source.get('source_type')} at {source.get('source_location')} -> {source.get('data_class')} ({source.get('last_status') or 'unknown'})",
            style="List Bullet",
        )

    if export_context.get("search_result"):
        document.add_heading("Redacted Search Results", level=1)
        search_result = export_context["search_result"]
        document.add_paragraph(f"Query: {search_result.get('query')}")
        document.add_paragraph(f"Patient hits: {len(search_result.get('patient_hits', []))}")
        document.add_paragraph(f"Study hits: {len(search_result.get('study_hits', []))}")
        document.add_paragraph(f"Asset hits: {len(search_result.get('asset_hits', []))}")

    if include_diagrams:
        document.add_heading("Diagrams", level=1)
        for diagram in diagrams:
            document.add_heading(diagram["name"].title(), level=2)
            document.add_picture(diagram["path"], width=Inches(6.8))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(output_path)


def _render_pdf(export_context, diagrams, output_path, include_diagrams=True):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallMono", parent=styles["BodyText"], fontName="Courier", fontSize=9, leading=11))

    story = [
        Paragraph("PACS Mentor Export", styles["Title"]),
        Spacer(1, 0.12 * inch),
        Paragraph("Read-only registry summary, workflow diagrams, and export-ready PACS metadata", styles["Italic"]),
        Spacer(1, 0.16 * inch),
    ]

    for line in export_context["summary_lines"]:
        story.append(Paragraph(line, styles["BodyText"]))

    story.extend([Spacer(1, 0.12 * inch), Paragraph("Registry Overview", styles["Heading1"]), Paragraph(export_context["overview"], styles["BodyText"])])

    snapshot = export_context["snapshot"]
    story.append(Spacer(1, 0.1 * inch))
    story.append(Paragraph("Source Inventory", styles["Heading1"]))
    for source in snapshot.get("sources", []):
        story.append(
            Paragraph(
                f"- {source.get('source_type')} at {source.get('source_location')} -> {source.get('data_class')} ({source.get('last_status') or 'unknown'})",
                styles["BodyText"],
            )
        )

    if export_context.get("search_result"):
        search_result = export_context["search_result"]
        story.append(Spacer(1, 0.1 * inch))
        story.append(Paragraph("Redacted Search Results", styles["Heading1"]))
        story.append(Paragraph(f"Query: {search_result.get('query')}", styles["BodyText"]))
        story.append(Paragraph(f"Patient hits: {len(search_result.get('patient_hits', []))}", styles["BodyText"]))
        story.append(Paragraph(f"Study hits: {len(search_result.get('study_hits', []))}", styles["BodyText"]))
        story.append(Paragraph(f"Asset hits: {len(search_result.get('asset_hits', []))}", styles["BodyText"]))

    if include_diagrams:
        story.append(PageBreak())
        story.append(Paragraph("Diagrams", styles["Heading1"]))
        for diagram in diagrams:
            story.append(Spacer(1, 0.08 * inch))
            story.append(Paragraph(diagram["name"].title(), styles["Heading2"]))
            story.append(RLImage(diagram["path"], width=6.8 * inch, height=3.8 * inch))

    doc = SimpleDocTemplate(output_path, pagesize=letter, leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.65 * inch, bottomMargin=0.65 * inch)
    doc.build(story)


def generate_pacs_export_artifact(registry, mentor=None, output_format="docx", query=None, redact=True, include_diagrams=True, output_dir=None):
    output_format = (output_format or "docx").lower()
    if output_format not in {"docx", "pdf"}:
        raise ValueError("output_format must be docx or pdf")

    output_root = output_dir or tempfile.mkdtemp(prefix="pacs_export_")
    os.makedirs(output_root, exist_ok=True)
    export_context = _build_export_context(registry, mentor=mentor, query=query, redact=redact)
    temp_dir = os.path.join(output_root, "diagrams")
    os.makedirs(temp_dir, exist_ok=True)
    diagrams = _create_diagram_set(temp_dir) if include_diagrams else []

    base_name = _safe_filename(f"pacs_export_{_now_stamp()}_{output_format}")
    output_path = os.path.join(output_root, f"{base_name}.{output_format}")

    if output_format == "docx":
        _render_docx(export_context, diagrams, output_path, include_diagrams=include_diagrams)
        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        _render_pdf(export_context, diagrams, output_path, include_diagrams=include_diagrams)
        mime_type = "application/pdf"

    return {
        "file_path": output_path,
        "file_name": os.path.basename(output_path),
        "mime_type": mime_type,
        "output_format": output_format,
        "export_context": export_context,
    }